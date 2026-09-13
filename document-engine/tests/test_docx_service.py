from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.services.docx import DocxService, substitute_placeholders
from app.services.templates import TemplateLoader

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"


@pytest.fixture()
def resolved_template():
    loader = TemplateLoader(TEMPLATES_DIR)
    return loader.resolve("solicitacao_aquisicao_ti", "1.0")


@pytest.fixture()
def sample_fields() -> dict[str, str]:
    return {
        "numero_solicitacao": "SA-2026-0042",
        "data_solicitacao": "2026-09-12",
        "solicitante": "João Silva",
        "area_solicitante": "TI - Infraestrutura",
        "fornecedor": "Dell Technologies",
        "valor": "R$ 45.000,00",
        "justificativa": "Substituição de equipamentos obsoletos da agência.",
        "categoria_produto": "Notebooks",
        "aprovacao_cade": "Não aplicável",
        "observacoes": "",
    }


class TestDocxService:
    def test_generate_creates_file(
        self, resolved_template, sample_fields, tmp_path: Path
    ):
        service = DocxService()
        output = tmp_path / "test_output.docx"
        result = service.generate(resolved_template, sample_fields, output, title="Test")
        assert result == output
        assert output.is_file()
        assert output.stat().st_size > 1000

    def test_generated_docx_is_valid(
        self, resolved_template, sample_fields, tmp_path: Path
    ):
        service = DocxService()
        output = tmp_path / "valid.docx"
        service.generate(resolved_template, sample_fields, output)
        doc = Document(str(output))
        assert len(doc.paragraphs) > 0

    def test_placeholders_substituted(
        self, resolved_template, sample_fields, tmp_path: Path
    ):
        service = DocxService()
        output = tmp_path / "substituted.docx"
        service.generate(resolved_template, sample_fields, output)

        doc = Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "João Silva" in full_text
        assert "Dell Technologies" in full_text
        assert "SA-2026-0042" in full_text
        assert "{{solicitante}}" not in full_text
        assert "{{fornecedor}}" not in full_text

    def test_no_residual_placeholders(
        self, resolved_template, sample_fields, tmp_path: Path
    ):
        service = DocxService()
        output = tmp_path / "no_residual.docx"
        service.generate(resolved_template, sample_fields, output)

        doc = Document(str(output))
        for paragraph in doc.paragraphs:
            assert "{{" not in paragraph.text or "observacoes" in paragraph.text, (
                f"Residual placeholder found: {paragraph.text!r}"
            )

    def test_title_set_in_core_properties(
        self, resolved_template, sample_fields, tmp_path: Path
    ):
        service = DocxService()
        output = tmp_path / "titled.docx"
        service.generate(resolved_template, sample_fields, output, title="My Title")

        doc = Document(str(output))
        assert doc.core_properties.title == "My Title"

    def test_template_not_modified(self, resolved_template, sample_fields, tmp_path: Path):
        original_size = resolved_template.docx_path.stat().st_size
        service = DocxService()
        output = tmp_path / "output.docx"
        service.generate(resolved_template, sample_fields, output)
        assert resolved_template.docx_path.stat().st_size == original_size

    def test_filename_safe(self, resolved_template, sample_fields, tmp_path: Path):
        service = DocxService()
        output = tmp_path / "sub dir" / "my file.docx"
        result = service.generate(resolved_template, sample_fields, output)
        assert result.is_file()


class TestSubstitutePlaceholders:
    def test_simple_replacement(self):
        doc = Document()
        doc.add_paragraph("Hello {{name}}, welcome!")
        remaining = substitute_placeholders(doc, {"name": "Alice"})
        assert remaining == []
        assert doc.paragraphs[0].text == "Hello Alice, welcome!"

    def test_multiple_placeholders(self):
        doc = Document()
        doc.add_paragraph("{{first}} and {{second}}")
        remaining = substitute_placeholders(doc, {"first": "A", "second": "B"})
        assert remaining == []
        assert doc.paragraphs[0].text == "A and B"

    def test_missing_field_leaves_placeholder(self):
        doc = Document()
        doc.add_paragraph("Value: {{missing_field}}")
        remaining = substitute_placeholders(doc, {})
        assert "missing_field" in remaining

    def test_empty_value_replaces_with_empty(self):
        doc = Document()
        doc.add_paragraph("Value: {{field}}")
        remaining = substitute_placeholders(doc, {"field": ""})
        assert remaining == []
        assert doc.paragraphs[0].text == "Value: "

    def test_special_characters_preserved(self):
        doc = Document()
        doc.add_paragraph("{{text}}")
        remaining = substitute_placeholders(doc, {"text": "R$ 1.000,00 <>&"})
        assert remaining == []
        assert "R$ 1.000,00 <>&" in doc.paragraphs[0].text
