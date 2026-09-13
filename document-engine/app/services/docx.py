from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn

from app.services.templates import ResolvedTemplate

logger = logging.getLogger(__name__)

_PLACEHOLDER_RE = re.compile(r"\{\{([a-zA-Z0-9_]+)\}\}")


def _iter_paragraphs(doc: Document) -> Iterable:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            if hdr is not None and hdr.is_linked_to_previous is False:
                yield from hdr.paragraphs
        for ftr in (section.footer, section.first_page_footer, section.even_page_footer):
            if ftr is not None and ftr.is_linked_to_previous is False:
                yield from ftr.paragraphs


def _substitute_in_paragraph(paragraph, fields: dict[str, str]) -> None:
    runs = list(paragraph.runs)
    if not runs:
        return

    texts = [r.text or "" for r in runs]
    merged = "".join(texts)

    if not _PLACEHOLDER_RE.search(merged):
        return

    offsets: list[tuple[int, int]] = []
    cursor = 0
    for t in texts:
        offsets.append((cursor, cursor + len(t)))
        cursor += len(t)

    def run_index_for(pos: int) -> int:
        for i, (s, e) in enumerate(offsets):
            if s <= pos < e:
                return i
        return len(runs) - 1

    matches = list(_PLACEHOLDER_RE.finditer(merged))
    if not matches:
        return

    processed: set[str] = set()

    for match in reversed(matches):
        placeholder = match.group(0)
        field_name = match.group(1)
        if placeholder in processed:
            continue
        processed.add(placeholder)

        if field_name not in fields:
            continue
        value = fields.get(field_name, "")
        if value is None:
            value = ""

        start, end = match.start(), match.end()
        first_idx = run_index_for(start)
        last_idx = run_index_for(max(start, end - 1))

        if first_idx == last_idx:
            run = runs[first_idx]
            run_text = run.text or ""
            local_start = start - offsets[first_idx][0]
            local_end = end - offsets[first_idx][0]
            run.text = run_text[:local_start] + str(value) + run_text[local_end:]
        else:
            first_run = runs[first_idx]
            first_text = first_run.text or ""
            local_start = start - offsets[first_idx][0]
            first_run.text = first_text[:local_start] + str(value)

            last_run = runs[last_idx]
            last_text = last_run.text or ""
            local_end = end - offsets[last_idx][0]
            last_run.text = last_text[local_end:]

            for i in range(first_idx + 1, last_idx):
                runs[i].text = ""


def substitute_placeholders(doc: Document, fields: dict[str, str]) -> list[str]:
    for paragraph in _iter_paragraphs(doc):
        _substitute_in_paragraph(paragraph, fields)

    return _check_remaining_placeholders(doc)


def _check_remaining_placeholders(doc: Document) -> list[str]:
    remaining: list[str] = []
    seen: set[str] = set()
    for paragraph in _iter_paragraphs(doc):
        for match in _PLACEHOLDER_RE.finditer(paragraph.text or ""):
            name = match.group(1)
            if name not in seen:
                seen.add(name)
                remaining.append(name)
    return remaining


class DocxService:
    def __init__(self) -> None:
        pass

    def generate(
        self,
        template: ResolvedTemplate,
        fields: dict[str, str],
        output_path: Path,
        title: str | None = None,
    ) -> Path:
        src = template.docx_path
        if not src.is_file():
            raise FileNotFoundError(f"Template not found: {src}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(src), str(output_path))

        doc = Document(str(output_path))

        remaining = substitute_placeholders(doc, fields)

        if title:
            core = doc.core_properties
            core.title = title

        doc.save(str(output_path))

        logger.info(
            "DOCX generated: file=%s template=%s v%s size=%d",
            output_path.name,
            template.key,
            template.version,
            output_path.stat().st_size,
        )

        if remaining:
            logger.warning(
                "Unsubstituted placeholders in output: %s",
                remaining,
            )

        return output_path
