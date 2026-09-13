"""Tests for memorando_interno and relatorio_operacional templates."""
from __future__ import annotations

import uuid
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.services.docx import DocxService, substitute_placeholders
from app.services.templates import TemplateLoader

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"


@pytest.fixture()
def loader() -> TemplateLoader:
    return TemplateLoader(TEMPLATES_DIR)


@pytest.fixture()
def memorando_template():
    loader = TemplateLoader(TEMPLATES_DIR)
    return loader.resolve("memorando_interno", "1.0")


@pytest.fixture()
def relatorio_template():
    loader = TemplateLoader(TEMPLATES_DIR)
    return loader.resolve("relatorio_operacional", "1.0")


@pytest.fixture()
def memorando_fields() -> dict[str, str]:
    return {
        "numero_documento": "MEM-2026-0001",
        "data": "2026-09-12",
        "destinatario": "Departamento de TI",
        "remetente": "João Silva",
        "departamento": "Recursos Humanos",
        "assunto": "Solicitação de Treinamento",
        "objetivo": "Capacitar equipe em novas tecnologias de desenvolvimento.",
        "conteudo": "Solicitamos a realização de treinamento em Python e FastAPI para a equipe de desenvolvimento. O treinamento deve cobrir conceitos básicos e avançados, com duração estimada de 40 horas.",
        "observacoes": "Preferência por treinamento presencial.",
    }


@pytest.fixture()
def relatorio_fields() -> dict[str, str]:
    return {
        "numero_relatorio": "REL-2026-0042",
        "data": "2026-09-12",
        "responsavel": "Maria Santos",
        "area": "Operações",
        "periodo": "Agosto 2026",
        "objetivo": "Registrar atividades operacionais e resultados do mês.",
        "resumo_executivo": "O mês de agosto apresentou resultados positivos com aumento de 15% na produtividade.",
        "atividades": "- Atendimento a 1.200 clientes\n- Processamento de 850 solicitações\n- Resolução de 45 incidentes",
        "resultados": "Meta de atendimento atingida em 110%. Satisfação do cliente em 92%.",
        "ocorrencias": "Duas interrupções de sistema registradas, ambas resolvidas em menos de 2 horas.",
        "recomendacoes": "Investir em automação de processos repetitivos.",
        "observacoes": "Equipe completa durante todo o período.",
    }


class TestMemorandoTemplate:
    def test_template_exists(self, memorando_template):
        assert memorando_template.docx_path.is_file()
        assert memorando_template.key == "memorando_interno"
        assert memorando_template.version == "1.0"

    def test_manifest_fields(self, memorando_template):
        assert "data" in memorando_template.manifest.required_fields
        assert "destinatario" in memorando_template.manifest.required_fields
        assert "remetente" in memorando_template.manifest.required_fields
        assert "observacoes" in memorando_template.manifest.optional_fields

    def test_generate_docx(self, memorando_template, memorando_fields, tmp_path: Path):
        service = DocxService()
        output = tmp_path / "memorando.docx"
        result = service.generate(memorando_template, memorando_fields, output)
        assert result.is_file()
        assert result.stat().st_size > 1000

    def test_placeholders_substituted(self, memorando_template, memorando_fields, tmp_path: Path):
        service = DocxService()
        output = tmp_path / "memorando_sub.docx"
        service.generate(memorando_template, memorando_fields, output)

        doc = Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Departamento de TI" in full_text
        assert "João Silva" in full_text
        assert "Recursos Humanos" in full_text
        assert "{{destinatario}}" not in full_text
        assert "{{remetente}}" not in full_text

    def test_missing_required_field(self, loader, memorando_template, tmp_path: Path):
        incomplete_fields = {
            "data": "2026-09-12",
            "destinatario": "TI",
            # Missing: remetente, departamento, assunto, objetivo, conteudo
        }
        missing = loader.validate_required_fields(memorando_template.manifest, incomplete_fields)
        assert len(missing) > 0
        assert "remetente" in missing
        assert "departamento" in missing


class TestRelatorioTemplate:
    def test_template_exists(self, relatorio_template):
        assert relatorio_template.docx_path.is_file()
        assert relatorio_template.key == "relatorio_operacional"
        assert relatorio_template.version == "1.0"

    def test_manifest_fields(self, relatorio_template):
        assert "data" in relatorio_template.manifest.required_fields
        assert "responsavel" in relatorio_template.manifest.required_fields
        assert "area" in relatorio_template.manifest.required_fields
        assert "resumo_executivo" in relatorio_template.manifest.required_fields
        assert "atividades" in relatorio_template.manifest.optional_fields
        assert "resultados" in relatorio_template.manifest.optional_fields

    def test_generate_docx(self, relatorio_template, relatorio_fields, tmp_path: Path):
        service = DocxService()
        output = tmp_path / "relatorio.docx"
        result = service.generate(relatorio_template, relatorio_fields, output)
        assert result.is_file()
        assert result.stat().st_size > 1000

    def test_placeholders_substituted(self, relatorio_template, relatorio_fields, tmp_path: Path):
        service = DocxService()
        output = tmp_path / "relatorio_sub.docx"
        service.generate(relatorio_template, relatorio_fields, output)

        doc = Document(str(output))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Maria Santos" in full_text
        assert "Operações" in full_text
        assert "Agosto 2026" in full_text
        assert "{{responsavel}}" not in full_text
        assert "{{area}}" not in full_text

    def test_missing_required_field(self, loader, relatorio_template, tmp_path: Path):
        incomplete_fields = {
            "data": "2026-09-12",
            "responsavel": "Ana",
            # Missing: area, periodo, objetivo, resumo_executivo
        }
        missing = loader.validate_required_fields(relatorio_template.manifest, incomplete_fields)
        assert len(missing) > 0
        assert "area" in missing
        assert "resumo_executivo" in missing


class TestTemplateSelection:
    def test_three_templates_available(self, loader):
        keys = loader.list_templates()
        assert "solicitacao_aquisicao_ti" in keys
        assert "memorando_interno" in keys
        assert "relatorio_operacional" in keys

    def test_exact_match_required(self, loader):
        # Should find exact match
        template = loader.resolve("memorando_interno", "1.0")
        assert template.key == "memorando_interno"

    def test_no_partial_match(self, loader):
        # Should NOT find partial match
        from app.services.templates import TemplateError
        with pytest.raises(TemplateError) as exc_info:
            loader.resolve("memorando", "1.0")  # Partial key should fail
        assert exc_info.value.code == "TEMPLATE_NOT_FOUND"

    def test_version_specific(self, loader):
        # Should find v1.0
        template = loader.resolve("relatorio_operacional", "1.0")
        assert template.version == "1.0"

        # Should NOT find v2.0 (doesn't exist)
        from app.services.templates import TemplateError
        with pytest.raises(TemplateError) as exc_info:
            loader.resolve("relatorio_operacional", "2.0")
        assert exc_info.value.code == "TEMPLATE_NOT_FOUND"


class TestGenerateEndpoint:
    def test_generate_memorando(self, client: TestClient, auth_headers: dict):
        spec = {
            "spec_id": str(uuid.uuid4()),
            "run_id": str(uuid.uuid4()),
            "task_id": str(uuid.uuid4()),
            "document_type": "memorando",
            "template_key": "memorando_interno",
            "template_version": "1.0",
            "title": "Memorando Interno",
            "fields": {
                "numero_documento": "MEM-2026-0001",
                "data": "2026-09-12",
                "destinatario": "Departamento de TI",
                "remetente": "João Silva",
                "departamento": "Recursos Humanos",
                "assunto": "Solicitação de Treinamento",
                "objetivo": "Capacitar equipe em novas tecnologias.",
                "conteudo": "Solicitamos treinamento em Python e FastAPI.",
                "observacoes": "",
            },
            "output": {"formats": ["docx"], "filename_prefix": "MEM-2026-0001"},
        }
        response = client.post("/generate", json=spec, headers=auth_headers)
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "generated"
        assert data["template_used"]["key"] == "memorando_interno"
        assert len(data["files"]) == 1
        assert data["files"][0]["format"] == "docx"

    def test_generate_relatorio(self, client: TestClient, auth_headers: dict):
        spec = {
            "spec_id": str(uuid.uuid4()),
            "run_id": str(uuid.uuid4()),
            "task_id": str(uuid.uuid4()),
            "document_type": "relatorio",
            "template_key": "relatorio_operacional",
            "template_version": "1.0",
            "title": "Relatório Operacional",
            "fields": {
                "numero_relatorio": "REL-2026-0042",
                "data": "2026-09-12",
                "responsavel": "Maria Santos",
                "area": "Operações",
                "periodo": "Agosto 2026",
                "objetivo": "Registrar atividades operacionais.",
                "resumo_executivo": "Resultados positivos com aumento de 15%.",
                "atividades": "Atendimento a 1.200 clientes.",
                "resultados": "Meta atingida em 110%.",
                "ocorrencias": "",
                "recomendacoes": "",
                "observacoes": "",
            },
            "output": {"formats": ["docx"], "filename_prefix": "REL-2026-0042"},
        }
        response = client.post("/generate", json=spec, headers=auth_headers)
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "generated"
        assert data["template_used"]["key"] == "relatorio_operacional"
        assert len(data["files"]) == 1
        assert data["files"][0]["format"] == "docx"

    def test_invalid_template_key_rejected(self, client: TestClient, auth_headers: dict):
        spec = {
            "spec_id": str(uuid.uuid4()),
            "run_id": str(uuid.uuid4()),
            "task_id": str(uuid.uuid4()),
            "document_type": "unknown",
            "template_key": "template_inexistente",  # Invalid key
            "template_version": "1.0",
            "title": "Test",
            "fields": {},
            "output": {"formats": ["docx"], "filename_prefix": "test"},
        }
        response = client.post("/generate", json=spec, headers=auth_headers)
        assert response.status_code == 404
        data = response.json()
        assert "TEMPLATE_NOT_FOUND" in str(data)

    def test_missing_required_field_rejected(self, client: TestClient, auth_headers: dict):
        spec = {
            "spec_id": str(uuid.uuid4()),
            "run_id": str(uuid.uuid4()),
            "task_id": str(uuid.uuid4()),
            "document_type": "memorando",
            "template_key": "memorando_interno",
            "template_version": "1.0",
            "title": "Memorando",
            "fields": {
                "data": "2026-09-12",
                "destinatario": "TI",
                # Missing required fields: remetente, departamento, assunto, objetivo, conteudo
            },
            "output": {"formats": ["docx"], "filename_prefix": "test"},
        }
        response = client.post("/generate", json=spec, headers=auth_headers)
        assert response.status_code == 400
        data = response.json()
        assert "MISSING_REQUIRED_FIELD" in str(data)
