"""Builds the memorando_interno template.docx from scratch using python-docx.

Run once to regenerate the template file:
    python scripts/build_memorando.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def build_memorando(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15

    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(1.18)

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run("ABIS — Agentic Banking Intelligence System")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        run.font.name = "Calibri"

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("Documento gerado pelo ABIS  •  Página ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        run.font.name = "Calibri"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("MEMORANDO INTERNO")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x0B, 0x30, 0x5E)

    ref = doc.add_paragraph()
    ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref.paragraph_format.space_after = Pt(18)
    run = ref.add_run("Nº {{numero_documento}}  •  {{data}}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    run.font.name = "Calibri"

    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(6)
    run = divider.add_run("━" * 72)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0x0B, 0x30, 0x5E)

    def section_heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x0B, 0x30, 0x5E)
        run.font.name = "Calibri"

    def field_row(label: str, placeholder: str, multi: bool = False) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.size = Pt(11)
        label_run.font.name = "Calibri"
        label_run.font.color.rgb = RGBColor(0x37, 0x47, 0x5E)
        value_run = p.add_run(placeholder)
        value_run.font.size = Pt(11)
        value_run.font.name = "Calibri"
        if multi:
            p.paragraph_format.space_after = Pt(8)

    section_heading("IDENTIFICAÇÃO")
    field_row("Para", "{{destinatario}}")
    field_row("De", "{{remetente}}")
    field_row("Departamento", "{{departamento}}")
    field_row("Data", "{{data}}")
    field_row("Assunto", "{{assunto}}")

    section_heading("OBJETIVO")
    field_row("Objetivo", "{{objetivo}}", multi=True)

    section_heading("CONTEÚDO")
    field_row("Conteúdo", "{{conteudo}}", multi=True)

    section_heading("OBSERVAÇÕES")
    field_row("Observações", "{{observacoes}}", multi=True)

    section_heading("REFERÊNCIAS NORMATIVAS")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(
        "Este memorando observa as normas internas aplicáveis conforme identificadas "
        "pelo ABIS no momento da emissão deste documento."
    )
    run.font.size = Pt(10)
    run.font.name = "Calibri"
    run.italic = True
    run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    sig_line = doc.add_paragraph()
    sig_line.paragraph_format.space_before = Pt(24)
    run = sig_line.add_run("_____________________________________________")
    run.font.size = Pt(10)
    run.font.name = "Calibri"

    sig_name = doc.add_paragraph()
    sig_name.paragraph_format.space_after = Pt(0)
    run = sig_name.add_run("{{remetente}}")
    run.font.size = Pt(10)
    run.font.name = "Calibri"

    sig_role = doc.add_paragraph()
    sig_role.paragraph_format.space_after = Pt(2)
    run = sig_role.add_run("{{departamento}}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    run.font.name = "Calibri"

    sig_date = doc.add_paragraph()
    sig_date.paragraph_format.space_after = Pt(2)
    run = sig_date.add_run("Data: {{data}}")
    run.font.size = Pt(10)
    run.font.name = "Calibri"

    doc.save(str(output_path))
    print(f"Memorando template written to: {output_path}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    target = here.parent / "templates" / "memorando_interno" / "v1.0" / "template.docx"
    build_memorando(target)
